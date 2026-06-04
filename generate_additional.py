"""
generate_additional.py
Adds 880 new documents to the existing synthetic_pdfs folder,
bringing the total corpus to ~1,000 documents.

Distribution is weighted to look like a real PE/wealth management firm.
Does NOT delete existing documents — only adds new ones.

Usage:
    python generate_additional.py
"""

import os
import random
from datetime import datetime, timedelta

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Pt

OUTPUT_DIR = "./synthetic_pdfs"
NUM_NEW    = 880

# ── WEIGHTED DISTRIBUTION ─────────────────────────────────────────────────────
# Reflects realistic PE/wealth management firm allocation
ASSET_WEIGHTS = {
    "PE":  220,
    "RE":  180,
    "PC":  160,
    "VC":  120,
    "SEC": 120,
    "OPP": 80,
}

# 60% Notes (Evoke-produced), 40% Marketing (third-party)
DOC_TYPE_WEIGHTS = {"Notes": 0.60, "Marketing": 0.40}

# Finance firms: mostly PDF and DOCX, some XLSX and PPTX
EXT_WEIGHTS = {"pdf": 0.40, "docx": 0.35, "xlsx": 0.15, "pptx": 0.10}

# Expanded manager lists per asset class
MANAGERS = {
    "PE": [
        "Blackstone", "CarlyleGroup", "KKR", "ApolloGlobal", "WarburgPincus",
        "TPGCapital", "AdventInternational", "BainCapital", "CVC Capital",
        "HelmanFriedman", "SilverLakePartners", "ThomasBravo", "VistaeEquity",
        "FranciscoPartners", "GoldenGate",
    ],
    "RE": [
        "Rialto", "BrookfieldRealEstate", "HinesAdvisors", "StarwoodCapital",
        "PrologisVentures", "BlackstoneRE", "CarlyleRealty", "AEWCapital",
        "AngelloGordon", "BentallGreenOak", "LaSalleInvestment", "NuveenRE",
        "OxfordProperties", "TishmanSpeyer", "RXRRealty",
    ],
    "PC": [
        "BlueMountain", "OakTreeCapital", "AresCreditMgmt", "GoldmanMezzanine",
        "AngeloDordon", "BlueBayAsset", "CarvalCredit", "GoldentreeAsset",
        "KingslandCapital", "MarathonAsset", "MidOceanCredit", "NinetyOne",
        "PGIMFixed", "SoundPointCapital", "TanagerCredit",
    ],
    "SEC": [
        "HarbourVest", "LexingtonPartners", "PantheonSecondaries", "StapletonGroup",
        "NorthseaFund", "ArdianSecondaries", "CovaBridge", "FollowOnCapital",
        "GICSecondaries", "HamiltonLane", "MotionEquity", "NewburyPartners",
        "SLCapital", "StepStone", "VertexSecondaries",
    ],
    "VC": [
        "AltitudeVentures", "SequoiaCapital", "a16zCrypto", "LightspeedVenture",
        "GeneralCatalyst", "AccelPartners", "BenchmarkCapital", "FirstRound",
        "FoundersFund", "GraycroftPartners", "IndexVentures", "KleinerPerkins",
        "NorwestVenture", "USVentures", "UnionSquareVentures",
    ],
    "OPP": [
        "BridgewaterAssociates", "TwosigmaInvestments", "ManInvestments",
        "WinctonGroup", "AQRCapital", "CambridgeAssociates", "DESHAWResearch",
        "ElementCapital", "GrahamCapital", "MilleniumMgmt",
    ],
}

# ── CONTENT POOLS ─────────────────────────────────────────────────────────────
NOTES_PARAS = [
    "Evoke Advisors internal analysis: The fund has demonstrated consistent performance across multiple market cycles, with particular strength in value-add strategies within the core plus segment.",
    "Portfolio review conducted by Evoke investment team. Current allocation reflects a defensive posture given prevailing macroeconomic headwinds including elevated interest rates and tightening credit conditions.",
    "Manager meeting notes — Evoke Advisors. Key discussion points included deployment pace, pipeline quality, and LP co-investment opportunities. Management expressed confidence in near-term exit environment.",
    "Evoke internal memo: Risk assessment for current vintage. Concentration risk remains elevated in top five positions. Recommend maintaining current hold period pending market stabilization.",
    "Quarterly monitoring report prepared by Evoke Advisors. Net IRR tracks above benchmark on a since-inception basis. TVPI of 1.4x reflects mark-to-market adjustments as of most recent valuation date.",
    "Due diligence summary — Evoke Advisors investment committee. Team quality rated above average. Strategy differentiation confirmed through reference checks with limited partners in prior funds.",
    "Evoke Advisors capital call notice received and processed. Unfunded commitment updated in portfolio management system. Pacing analysis reflects continued overweight to private credit.",
    "Annual meeting notes prepared by Evoke Advisors. Management presented updated five-year deployment plan and discussed macro headwinds affecting exit timing. Q&A session covered co-investment pipeline.",
    "Evoke Advisors monitoring update: Fund performance in line with expectations. Top three portfolio companies contributing 68% of total value. Watch list contains two positions under enhanced monitoring.",
    "Investment committee memo — Evoke Advisors. Recommending approval of follow-on investment subject to updated terms review. Board representation and information rights confirmed satisfactory.",
    "Evoke Advisors: Distribution notice received. Net proceeds allocated per waterfall calculation. LPA provisions reviewed and confirmed compliant with distribution policy.",
    "Fund review — Evoke Advisors. Vintage year analysis shows above-median performance relative to Cambridge Associates benchmark. Attribution analysis credits operational value creation program.",
]

MARKETING_PARAS = [
    "This document has been prepared by the manager for distribution to qualified institutional investors. Past performance is not indicative of future results. This material is confidential and proprietary.",
    "Investment highlights: The fund targets risk-adjusted returns in excess of public market equivalents through disciplined underwriting and active asset management across the capital structure.",
    "Manager overview: Founded in 2003, the firm manages over $45 billion in assets across flagship and co-investment vehicles. The team comprises 120 investment professionals across 8 global offices.",
    "Strategy overview: The fund pursues a concentrated, high-conviction approach focusing on mid-market companies with identifiable operational improvement opportunities and clear exit pathways.",
    "Performance summary as of most recent quarter-end. Net returns presented after management fees and carried interest. Benchmark comparison uses Cambridge Associates US Private Equity Index.",
    "Risk factors: Investments in private funds involve significant risks including illiquidity, loss of capital, and limited transparency. Prospective investors should carefully review all offering documents.",
    "Fee structure: 1.5% management fee on committed capital during the investment period, stepping down to 1.0% on invested capital thereafter. 20% carried interest above an 8% preferred return.",
    "Portfolio construction: The fund maintains a diversified portfolio of 15-20 investments across target sectors. Average hold period of 4-6 years with multiple exit pathways including strategic sale and IPO.",
    "ESG policy: The manager integrates environmental, social, and governance considerations into the investment process. Annual ESG reporting provided to limited partners on portfolio company metrics.",
    "Fund terms: Target fund size $2.5 billion with hard cap of $3.0 billion. Investment period of 5 years from final close. Fund life of 10 years with two one-year extension options.",
    "Co-investment program: Limited partners have access to co-investment opportunities alongside the main fund. Co-investments are offered on a no-fee, no-carry basis subject to capacity and eligibility.",
    "Track record: The manager has invested across three prior funds totaling $8.2 billion in commitments. Realized returns on fully exited investments show gross multiple of 2.8x and gross IRR of 24%.",
]

# ── HELPERS ───────────────────────────────────────────────────────────────────
def random_date():
    start = datetime(2019, 1, 1)
    end   = datetime(2025, 6, 30)
    return (start + timedelta(days=random.randint(0, (end - start).days))).strftime("%Y%m%d")

def weighted_choice(weight_dict):
    keys   = list(weight_dict.keys())
    values = list(weight_dict.values())
    total  = sum(values)
    probs  = [v / total for v in values]
    return random.choices(keys, weights=probs, k=1)[0]

def random_paras(doc_type, n=4):
    pool = NOTES_PARAS if doc_type == "Notes" else MARKETING_PARAS
    return random.sample(pool, min(n, len(pool)))

# ── FILE CREATORS ─────────────────────────────────────────────────────────────
def make_pdf(path, doc_type, asset, manager, date_str):
    doc  = fitz.open()
    page = doc.new_page()
    header = "EVOKE ADVISORS — INTERNAL" if doc_type == "Notes" else f"{manager.upper()} — INVESTOR MATERIALS"
    page.insert_text((50, 50), header, fontsize=11, fontname="helv")
    page.insert_text((50, 70), f"Asset: {asset}  |  Manager: {manager}  |  Date: {date_str}", fontsize=9, fontname="helv")
    y = 100
    for para in random_paras(doc_type, random.randint(3, 6)):
        words, line = para.split(), ""
        for word in words:
            if len(line) + len(word) + 1 > 90:
                page.insert_text((50, y), line, fontsize=9, fontname="helv")
                y += 14; line = word
            else:
                line = (line + " " + word).strip()
        if line:
            page.insert_text((50, y), line, fontsize=9, fontname="helv")
        y += 22
        if y > 700:
            page = doc.new_page()
            y = 50
    doc.save(path); doc.close()

def make_docx(path, doc_type, asset, manager, date_str):
    doc = Document()
    title = "EVOKE ADVISORS — INTERNAL" if doc_type == "Notes" else f"{manager.upper()} — INVESTOR MATERIALS"
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Asset Class: {asset}  |  Manager: {manager}  |  Date: {date_str}")
    doc.add_paragraph("")
    for para in random_paras(doc_type, random.randint(3, 6)):
        doc.add_paragraph(para)
    doc.save(path)

def make_xlsx(path, doc_type, asset, manager, date_str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fund Data"
    ws.append(["Asset Class", "Manager", "Doc Type", "Date", "Metric", "Value"])
    metrics = ["Net IRR", "TVPI", "DPI", "RVPI", "Commitment ($M)", "Called ($M)",
               "Distributed ($M)", "Remaining Value ($M)", "Total Value ($M)"]
    for metric in metrics:
        ws.append([asset, manager, doc_type, date_str, metric, round(random.uniform(0.5, 3.5), 2)])
    ws.append([])
    ws.append(["Notes:"])
    for para in random_paras(doc_type, 2):
        ws.append([para])
    wb.save(path)

def make_pptx(path, doc_type, asset, manager, date_str):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = "EVOKE ADVISORS — INTERNAL" if doc_type == "Notes" else f"{manager} — INVESTOR MATERIALS"
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"{asset}  |  {manager}  |  {date_str}"
    for para in random_paras(doc_type, random.randint(3, 5)):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{manager} — {doc_type}"
        slide.placeholders[1].text = para
    prs.save(path)

MAKERS = {"pdf": make_pdf, "docx": make_docx, "xlsx": make_xlsx, "pptx": make_pptx}

# ── MAIN ─────────────────────────────────────────────────────────────────────
def generate():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # Get existing filenames to avoid collisions
    existing = set(os.listdir(OUTPUT_DIR))
    name_counts = {}

    # Pre-populate name_counts from existing files to handle duplicates correctly
    for f in existing:
        import re
        m = re.match(r'^(.+?)( \(\d+\))?\.(pdf|docx|xlsx|pptx)$', f, re.IGNORECASE)
        if m:
            base_ext = f"{m.group(1)}.{m.group(3).lower()}"
            name_counts[base_ext] = name_counts.get(base_ext, 0) + 1

    # Build weighted asset list
    asset_pool = []
    for asset, count in ASSET_WEIGHTS.items():
        asset_pool.extend([asset] * count)
    random.shuffle(asset_pool)

    generated = 0
    idx       = 0

    print(f"\nGenerating {NUM_NEW} additional documents into {OUTPUT_DIR}/\n")

    while generated < NUM_NEW:
        asset  = asset_pool[idx % len(asset_pool)]
        idx   += 1
        mgr    = random.choice(MANAGERS[asset])
        dtype  = weighted_choice(DOC_TYPE_WEIGHTS)
        date   = random_date()
        ext    = weighted_choice(EXT_WEIGHTS)

        base_ext = f"{asset}.{mgr}.{dtype}.{date}.{ext}"
        count    = name_counts.get(base_ext, 0)
        filename = f"{asset}.{mgr}.{dtype}.{date}.{ext}" if count == 0 \
                   else f"{asset}.{mgr}.{dtype}.{date} ({count}).{ext}"
        name_counts[base_ext] = count + 1

        out_path = os.path.join(OUTPUT_DIR, filename)
        try:
            MAKERS[ext](out_path, dtype, asset, mgr, date)
            print(f"  [{generated+1:04d}] Created: {filename}")
            generated += 1
        except Exception as e:
            print(f"  ERROR: {filename} — {e}")

    print(f"\n✅  Done. {generated} new documents added to {OUTPUT_DIR}/")
    print(f"    Total corpus size: ~{len(os.listdir(OUTPUT_DIR))} documents")
    print(f"\n    Run next:")
    print(f"      python validate_filenames.py --dir ./synthetic_pdfs")
    print(f"      python phase1_ingest.py --dir ./synthetic_pdfs")
    print(f"      python phase2_chunk.py --dir ./synthetic_pdfs")
    print(f"      python phase3_embed.py")

if __name__ == "__main__":
    generate()
