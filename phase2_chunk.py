"""
phase2_chunk.py
Opens every document, extracts text, splits into semantic chunks,
and stores chunks in a new Supabase table ready for embedding.

Handles: PDF, DOCX, XLSX, PPTX
Each chunk carries: document_id, page_number, chunk_index, text, doc metadata

Usage:
    python phase2_chunk.py --dir ./synthetic_pdfs
    python phase2_chunk.py --dir ./real_docs    # on Wednesday
"""

import os
import re
import argparse
import psycopg2
from dotenv import load_dotenv

load_dotenv()

DATABASE_URL = os.getenv("DATABASE_URL")

# ── SCHEMA ────────────────────────────────────────────────────────────────────
CREATE_CHUNKS_TABLE = """
CREATE TABLE IF NOT EXISTS chunks (
    id              SERIAL PRIMARY KEY,
    chunk_id        TEXT UNIQUE NOT NULL,
    document_id     TEXT NOT NULL REFERENCES documents(document_id),
    filename        TEXT NOT NULL,
    manager_name    TEXT NOT NULL,
    asset_class     TEXT NOT NULL,
    doc_type        TEXT NOT NULL,
    doc_date        DATE NOT NULL,
    file_extension  TEXT NOT NULL,
    chunk_index     INTEGER NOT NULL,
    page_number     INTEGER,
    text            TEXT NOT NULL,
    word_count      INTEGER,
    created_at      TIMESTAMP DEFAULT NOW()
);
CREATE INDEX IF NOT EXISTS idx_chunk_doc_id    ON chunks(document_id);
CREATE INDEX IF NOT EXISTS idx_chunk_manager   ON chunks(manager_name);
CREATE INDEX IF NOT EXISTS idx_chunk_asset     ON chunks(asset_class);
CREATE INDEX IF NOT EXISTS idx_chunk_type      ON chunks(doc_type);
CREATE INDEX IF NOT EXISTS idx_chunk_date      ON chunks(doc_date);
"""

INSERT_CHUNK_SQL = """
INSERT INTO chunks (
    chunk_id, document_id, filename, manager_name, asset_class,
    doc_type, doc_date, file_extension, chunk_index, page_number,
    text, word_count
)
VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
ON CONFLICT (chunk_id) DO UPDATE SET
    text       = EXCLUDED.text,
    word_count = EXCLUDED.word_count;
"""

# ── TEXT EXTRACTION ───────────────────────────────────────────────────────────

def extract_pdf(file_path):
    """Returns list of (page_number, text) tuples."""
    import fitz
    pages = []
    doc = fitz.open(file_path)
    for page_num in range(doc.page_count):
        text = doc[page_num].get_text().strip()
        if text:
            pages.append((page_num + 1, text))
    doc.close()
    return pages

def extract_docx(file_path):
    """Returns list of (page_number, text) tuples.
    Word docs don't have true pages — we group every 5 paragraphs as one 'page'."""
    from docx import Document
    doc      = Document(file_path)
    paras    = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    pages    = []
    group_size = 5
    for i in range(0, len(paras), group_size):
        group     = paras[i:i + group_size]
        text      = " ".join(group)
        page_num  = (i // group_size) + 1
        if text:
            pages.append((page_num, text))
    return pages

def extract_xlsx(file_path):
    """Returns list of (sheet_number, text) tuples — one per sheet."""
    from openpyxl import load_workbook
    wb     = load_workbook(file_path, read_only=True, data_only=True)
    pages  = []
    for sheet_num, sheet_name in enumerate(wb.sheetnames, 1):
        ws    = wb[sheet_name]
        cells = []
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    cells.append(str(cell))
        text = " | ".join(cells)
        if text.strip():
            pages.append((sheet_num, f"Sheet: {sheet_name}\n{text}"))
    wb.close()
    return pages

def extract_pptx(file_path):
    """Returns list of (slide_number, text) tuples — one per slide."""
    from pptx import Presentation
    prs   = Presentation(file_path)
    pages = []
    for slide_num, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        text = " ".join(texts)
        if text:
            pages.append((slide_num, text))
    return pages

EXTRACTORS = {
    ".pdf":  extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
}

# ── CHUNKING ──────────────────────────────────────────────────────────────────

def chunk_text(text, max_words=400, overlap_words=40):
    """
    Splits text into overlapping chunks of ~max_words words.
    Tries to split on sentence boundaries for cleaner chunks.
    Returns list of text strings.
    """
    # Split into sentences
    sentences = re.split(r'(?<=[.!?])\s+', text.strip())
    sentences = [s.strip() for s in sentences if s.strip()]

    if not sentences:
        return []

    chunks  = []
    current = []
    current_words = 0

    for sentence in sentences:
        sentence_words = len(sentence.split())

        # If adding this sentence exceeds max, save current chunk and start new
        if current_words + sentence_words > max_words and current:
            chunk_text_str = " ".join(current)
            if chunk_text_str.strip():
                chunks.append(chunk_text_str)

            # Keep last few sentences as overlap for context continuity
            overlap = []
            overlap_count = 0
            for s in reversed(current):
                w = len(s.split())
                if overlap_count + w <= overlap_words:
                    overlap.insert(0, s)
                    overlap_count += w
                else:
                    break
            current       = overlap + [sentence]
            current_words = sum(len(s.split()) for s in current)
        else:
            current.append(sentence)
            current_words += sentence_words

    # Add final chunk
    if current:
        chunk_text_str = " ".join(current)
        if chunk_text_str.strip():
            chunks.append(chunk_text_str)

    return chunks

# ── MAIN ─────────────────────────────────────────────────────────────────────

def run_chunking(docs_dir):
    if not DATABASE_URL:
        raise EnvironmentError("DATABASE_URL not set in .env")

    print("\nConnecting to database...")
    conn = psycopg2.connect(DATABASE_URL)
    cur  = conn.cursor()

    print("Creating chunks table if not exists...")
    cur.execute(CREATE_CHUNKS_TABLE)
    conn.commit()

    # Load all document metadata from Supabase
    cur.execute("""
        SELECT document_id, filename, manager_name, asset_class,
               doc_type, doc_date, file_extension, file_path
        FROM documents
        ORDER BY filename
    """)
    docs = cur.fetchall()

    if not docs:
        print("No documents found in database. Run phase1_ingest.py first.")
        return

    print(f"\nChunking {len(docs)} documents...\n")

    total_chunks = 0
    errors       = []

    for i, (doc_id, filename, manager, asset, dtype, doc_date,
            ext, file_path) in enumerate(docs, 1):

        # Use the path stored in DB, but also try the docs_dir in case
        # files were moved
        if not os.path.exists(file_path):
            alt_path = os.path.join(docs_dir, filename)
            if os.path.exists(alt_path):
                file_path = alt_path
            else:
                print(f"  [{i:04d}] MISSING  {filename}")
                errors.append(filename)
                continue

        extractor = EXTRACTORS.get(ext.lower())
        if not extractor:
            print(f"  [{i:04d}] SKIP     {filename}  (unsupported type: {ext})")
            continue

        try:
            pages = extractor(file_path)
        except Exception as e:
            print(f"  [{i:04d}] ERROR    {filename}  →  {e}")
            errors.append(filename)
            continue

        doc_chunks = 0
        chunk_index = 0

        for page_num, page_text in pages:
            page_chunks = chunk_text(page_text)

            for chunk_str in page_chunks:
                chunk_id   = f"{doc_id}_chunk_{chunk_index:04d}"
                word_count = len(chunk_str.split())

                try:
                    cur.execute(INSERT_CHUNK_SQL, (
                        chunk_id,
                        doc_id,
                        filename,
                        manager,
                        asset,
                        dtype,
                        doc_date,
                        ext,
                        chunk_index,
                        page_num,
                        chunk_str,
                        word_count,
                    ))
                    chunk_index += 1
                    doc_chunks  += 1
                except Exception as e:
                    conn.rollback()
                    print(f"    Chunk insert error: {e}")
                    continue

        conn.commit()
        total_chunks += doc_chunks
        print(f"  [{i:04d}] OK  {filename}  →  {doc_chunks} chunks")

    cur.close()
    conn.close()

    # ── FINAL REPORT ──────────────────────────────────────────────────────
    print(f"\n{'=' * 70}")
    print("CHUNKING COMPLETE")
    print(f"{'=' * 70}")
    print(f"  Documents processed : {len(docs) - len(errors)}")
    print(f"  Total chunks created: {total_chunks}")
    print(f"  Avg chunks per doc  : {total_chunks // max(1, len(docs) - len(errors))}")
    print(f"  Errors              : {len(errors)}")

    if errors:
        print(f"\n  Failed files:")
        for f in errors:
            print(f"    {f}")

    if not errors:
        print(f"\n  ✅  All documents chunked successfully.")
        print(f"      Run next: python phase3_embed.py")
    print(f"{'=' * 70}\n")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dir", required=True, help="Path to documents directory")
    args = parser.parse_args()
    run_chunking(args.dir)
