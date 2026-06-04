"""
generate_synthetic.py
Generates synthetic documents (PDF, DOCX, XLSX, PPTX) mirroring the
exact Evoke Advisors naming convention.
"""

import os
import random
from datetime import datetime, timedelta

import fitz                          # PyMuPDF  — PDF
from docx import Document            # python-docx — Word
from openpyxl import Workbook        # openpyxl  — Excel
from pptx import Presentation        # python-pptx — PowerPoint
from pptx.util import Inches, Pt

OUTPUT_DIR  = "./synthetic_pdfs"
NUM_DOCS    = 120   # ~30 per file type

ASSET_CLASSES = ["RE", "PE", "PC", "SEC", "VC", "OPP"]

MANAGERS = {
    "RE":  ["Rialto", "BrookfieldRealEstate", "HinesAdvisors", "StarwoodCapital", "PrologisVentures"],
    "PE":  ["Blackstone", "CarlyleGroup", "KKR", "ApolloGlobal", "WarburgPincus"],
    "PC":  ["BlueMountain", "OakTreeCapital", "AresCreditMgmt", "GoldmanMezzanine", "AngeloDordon"],
    "SEC": ["HarbourVest", "LexingtonPartners", "PantheonSecondaries", "StapletonGroup", "NorthseaFund"],
    "VC":  ["AltitudeVentures", "SequoiaCapital", "a16zCrypto", "LightspeedVenture", "GeneralCatalyst"],
    "OPP": ["BridgewaterAssociates", "TwosigmaInvestments", "ManInvestments", "WinctonGroup", "AQRCapital"],
}

DOC_TYPES  = ["Notes", "Marketing"]
EXTENSIONS = ["pdf", "docx", "xlsx", "pptx"]

NOTES_PARAS = [
    "Evoke Advisors internal analysis: The fund has demonstrated consistent performance across multiple market cycles, with particular strength in value-add strategies.",
    "Portfolio review conducted by Evoke investment team. Current allocation reflects a defensive posture given prevailing macroeconomic headwinds including elevated interest rates.",
    "Manager meeting notes — Evoke Advisors. Key discussion points included deployment pace, pipeline quality, and LP co-investment opportunities.",
    "Evoke internal memo: Risk assessment for current vintage. Concentration risk remains elevated in top five positions. Recommend maintaining current hold period.",
    "Quarterly monitoring report prepared by Evoke Advisors. Net IRR tracks above benchmark on a since-inception basis. TVPI of 1.4x reflects mark-to-market adjustments.",
    "Due diligence summary — Evoke Advisors investment committee. Team quality rated above average. Strategy differentiation confirmed through reference checks.",
]

MARKETING_PARAS = [
    "This document has been prepared by the manager for distribution to qualified institutional investors. Past performance is not indicative of future results.",
    "Investment highlights: The fund targets risk-adjusted returns in excess of public market equivalents through disciplined underwriting and active asset management.",
    "Manager overview: Founded in 2003, the firm manages over $45 billion in assets across flagship and co-investment vehicles across 8 global offices.",
    "Strategy overview: The fund pursues a concentrated, high-conviction approach focusing on mid-market companies with identifiable operational improvement opportunities.",
    "Performance summary as of most recent quarter-end. Net returns presented after management fees and carried interest. Benchmark uses Cambridge Associates Index.",
    "Risk factors: Investments in private funds involve significant risks including illiquidity, loss of capital, and limited transparency.",
    "Fee structure: 1.5% management fee on committed capital during the investment period. 20% carried interest above an 8% preferred return.",
]

def random_date():
    start = datetime(2018, 1, 1)
    end   = datetime(2025, 6, 30)
    return (start + timedelta(days=random.randint(0, (end - start).days))).strftime("%Y%m%d")

def random_paras(doc_type, n=4):
    pool = NOTES_PARAS if doc_type == "Notes" else MARKETING_PARAS
    return random.sample(pool, min(n, len(pool)))

# ── FILE CREATORS ─────────────────────────────────────────────────────────────

def make_pdf(path, doc_type, asset, manager, date_str):
    doc  = fitz.open()
    page = doc.new_page()
    header = f"{'EVOKE ADVISORS — INTERNAL' if doc_type == 'Notes' else manager.upper() + ' — INVESTOR MATERIALS'}"
    page.insert_text((50, 50), header, fontsize=11, fontname="helv")
    page.insert_text((50, 70), f"Asset: {asset}  |  Manager: {manager}  |  Date: {date_str}", fontsize=9, fontname="helv")
    y = 100
    for para in random_paras(doc_type):
        words, line = para.split(), ""
        for word in words:
            if len(line) + len(word) + 1 > 90:
                page.insert_text((50, y), line, fontsize=9, fontname="helv")
                y += 14; line = word
            else:
                line = (line + " " + word).strip()
        if line:
            page.insert_text((50, y), line, fontsize=9, fontname="helv")
        y += 22
    doc.save(path); doc.close()

def make_docx(path, doc_type, asset, manager, date_str):
    doc = Document()
    title = "EVOKE ADVISORS — INTERNAL" if doc_type == "Notes" else f"{manager.upper()} — INVESTOR MATERIALS"
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Asset Class: {asset}  |  Manager: {manager}  |  Date: {date_str}")
    doc.add_paragraph("")
    for para in random_paras(doc_type):
        doc.add_paragraph(para)
    doc.save(path)

def make_xlsx(path, doc_type, asset, manager, date_str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fund Data"
    ws.append(["Asset Class", "Manager", "Doc Type", "Date", "Metric", "Value"])
    metrics = ["Net IRR", "TVPI", "DPI", "RVPI", "Commitment ($M)", "Called ($M)", "Distributed ($M)"]
    for i, metric in enumerate(metrics):
        ws.append([asset, manager, doc_type, date_str, metric, round(random.uniform(0.5, 3.5), 2)])
    ws.append([])
    ws.append(["Notes:"])
    for para in random_paras(doc_type, 2):
        ws.append([para])
    wb.save(path)

def make_pptx(path, doc_type, asset, manager, date_str):
    prs = Presentation()
    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = "EVOKE ADVISORS — INTERNAL" if doc_type == "Notes" else f"{manager} — INVESTOR MATERIALS"
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"{asset}  |  {manager}  |  {date_str}"
    # Content slides
    for para in random_paras(doc_type, 3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{manager} — {doc_type}"
        slide.placeholders[1].text = para
    prs.save(path)

MAKERS = {
    "pdf":  make_pdf,
    "docx": make_docx,
    "xlsx": make_xlsx,
    "pptx": make_pptx,
}

# ── MAIN ─────────────────────────────────────────────────────────────────────

def generate():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    name_counts = {}
    generated   = 0

    while generated < NUM_DOCS:
        asset  = random.choice(ASSET_CLASSES)
        mgr    = random.choice(MANAGERS[asset])
        dtype  = random.choice(DOC_TYPES)
        date   = random_date()
        ext    = random.choice(EXTENSIONS)
        base   = f"{asset}.{mgr}.{dtype}.{date}"

        count = name_counts.get(f"{base}.{ext}", 0)
        filename = f"{base}.{ext}" if count == 0 else f"{base} ({count}).{ext}"
        name_counts[f"{base}.{ext}"] = count + 1

        out_path = os.path.join(OUTPUT_DIR, filename)
        MAKERS[ext](out_path, dtype, asset, mgr, date)
        print(f"  Created: {filename}")
        generated += 1

    print(f"\nDone. {generated} documents written to {OUTPUT_DIR}/")

if __name__ == "__main__":
    generate()
