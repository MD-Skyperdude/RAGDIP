"""
phase1_ingest.py
Parses every document filename and writes one metadata row per file
into the Evoke PostgreSQL database. Handles PDF, DOCX, XLSX, PPTX.

Run AFTER validate_filenames.py shows zero invalid files.

Usage:
    python phase1_ingest.py --dir ./synthetic_pdfs
    python phase1_ingest.py --dir ./real_docs        # on Wednesday
"""

import os
import re
import uuid
import argparse
import psycopg2
from datetime import datetime
from dotenv import load_dotenv

load_dotenv()

DATABASE_URL        = os.getenv("DATABASE_URL")
VALID_ASSET_CLASSES = {"RE", "PE", "PC", "SEC", "VC", "OPP"}
VALID_DOC_TYPES     = {"Notes", "Marketing"}
VALID_EXTENSIONS    = {".pdf", ".docx", ".xlsx", ".pptx"}

ASSET_CLASS_LABELS = {
    "RE":  "Real Estate (Real Assets)",
    "PE":  "Private Equity",
    "PC":  "Private Credit",
    "SEC": "Secondaries",
    "VC":  "Venture Capital",
    "OPP": "Uncorrelated",
}

PATTERN = re.compile(
    r'^([A-Z]+)\.'
    r'([A-Za-z0-9]+)\.'
    r'(Notes|Marketing)\.'
    r'(\d{8})'
    r'(?: \((\d+)\))?'
    r'\.(pdf|docx|xlsx|pptx)$',
    re.IGNORECASE
)

CREATE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS documents (
    id              SERIAL PRIMARY KEY,
    document_id     TEXT UNIQUE NOT NULL,
    filename        TEXT NOT NULL,
    manager_name    TEXT NOT NULL,
    asset_class     TEXT NOT NULL,
    asset_label     TEXT NOT NULL,
    doc_type        TEXT NOT NULL,
    doc_date        DATE NOT NULL,
    duplicate_n     INTEGER,
    base_name       TEXT NOT NULL,
    file_path       TEXT NOT NULL,
    file_extension  TEXT NOT NULL,
    page_count      INTEGER,
    created_at      TIMESTAMP DEFAULT NOW()
);
CREATE INDEX IF NOT EXISTS idx_manager    ON documents(manager_name);
CREATE INDEX IF NOT EXISTS idx_asset      ON documents(asset_class);
CREATE INDEX IF NOT EXISTS idx_type       ON documents(doc_type);
CREATE INDEX IF NOT EXISTS idx_date       ON documents(doc_date);
CREATE INDEX IF NOT EXISTS idx_base_name  ON documents(base_name);
CREATE INDEX IF NOT EXISTS idx_extension  ON documents(file_extension);
"""

INSERT_SQL = """
INSERT INTO documents (
    document_id, filename, manager_name, asset_class, asset_label,
    doc_type, doc_date, duplicate_n, base_name, file_path,
    file_extension, page_count
)
VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
ON CONFLICT (document_id) DO UPDATE SET
    filename       = EXCLUDED.filename,
    file_path      = EXCLUDED.file_path,
    file_extension = EXCLUDED.file_extension,
    page_count     = EXCLUDED.page_count;
"""

def parse_filename(filename):
    m = PATTERN.match(filename)
    if not m:
        return None
    asset_class  = m.group(1).upper()
    manager_name = m.group(2)
    doc_type     = m.group(3).capitalize()
    date_str     = m.group(4)
    duplicate_n  = m.group(5)
    extension    = "." + m.group(6).lower()

    if asset_class not in VALID_ASSET_CLASSES:
        return None
    if doc_type not in VALID_DOC_TYPES:
        return None
    try:
        doc_date = datetime.strptime(date_str, "%Y%m%d").date()
    except ValueError:
        return None

    return {
        "asset_class":  asset_class,
        "asset_label":  ASSET_CLASS_LABELS[asset_class],
        "manager_name": manager_name,
        "doc_type":     doc_type,
        "doc_date":     doc_date,
        "extension":    extension,
        "duplicate_n":  int(duplicate_n) if duplicate_n else None,
        "base_name":    f"{asset_class}.{manager_name}.{doc_type}.{date_str}",
    }

def get_page_count(file_path, extension):
    try:
        if extension == ".pdf":
            import fitz
            doc = fitz.open(file_path)
            count = doc.page_count
            doc.close()
            return count
        elif extension == ".docx":
            # Word doesn't have pages natively — count paragraphs as proxy
            from docx import Document
            doc = Document(file_path)
            return max(1, len(doc.paragraphs) // 5)
        elif extension == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True)
            count = len(wb.sheetnames)
            wb.close()
            return count
        elif extension == ".pptx":
            from pptx import Presentation
            prs = Presentation(file_path)
            return len(prs.slides)
    except Exception:
        return None

def ingest(docs_dir):
    if not DATABASE_URL:
        raise EnvironmentError("DATABASE_URL not set in .env file")

    all_files = os.listdir(docs_dir)
    files = sorted([f for f in all_files
                    if os.path.splitext(f)[1].lower() in VALID_EXTENSIONS])

    if not files:
        print(f"No supported documents found in {docs_dir}")
        return

    print(f"\nConnecting to database...")
    conn = psycopg2.connect(DATABASE_URL)
    cur  = conn.cursor()
    print("Creating/verifying table schema...")
    cur.execute(CREATE_TABLE_SQL)
    conn.commit()

    print(f"\nIngesting {len(files)} documents...\n")

    success = 0
    skipped = 0
    failed  = []

    for i, filename in enumerate(files, 1):
        parsed = parse_filename(filename)
        if not parsed:
            print(f"  [{i:04d}] SKIP  {filename}")
            skipped += 1
            continue

        file_path   = os.path.abspath(os.path.join(docs_dir, filename))
        page_count  = get_page_count(file_path, parsed["extension"])
        document_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, filename))

        try:
            cur.execute(INSERT_SQL, (
                document_id,
                filename,
                parsed["manager_name"],
                parsed["asset_class"],
                parsed["asset_label"],
                parsed["doc_type"],
                parsed["doc_date"],
                parsed["duplicate_n"],
                parsed["base_name"],
                file_path,
                parsed["extension"],
                page_count,
            ))
            conn.commit()
            ext_label = parsed["extension"].upper().replace(".", "")
            print(f"  [{i:04d}] OK    {filename}  ({ext_label}, {page_count}p)")
            success += 1
        except Exception as e:
            conn.rollback()
            print(f"  [{i:04d}] ERROR {filename}  →  {e}")
            failed.append({"filename": filename, "error": str(e)})

    cur.close()
    conn.close()

    print(f"\n{'=' * 70}")
    print("INGESTION COMPLETE")
    print(f"{'=' * 70}")
    print(f"  Total files        : {len(files)}")
    print(f"  Inserted/updated   : {success}")
    print(f"  Skipped (bad name) : {skipped}")
    print(f"  Errors             : {len(failed)}")

    if failed:
        print(f"\n  Failed files:")
        for f in failed:
            print(f"    {f['filename']}  →  {f['error']}")

    if success > 0 and not failed:
        print(f"\n  ✅  All files ingested successfully.")
        print(f"      Verify in Supabase, then run: python phase2_chunk.py --dir {docs_dir}")
    print(f"{'=' * 70}\n")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dir", required=True, help="Path to documents directory")
    args = parser.parse_args()
    ingest(args.dir)
